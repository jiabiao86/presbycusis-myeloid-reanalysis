#!/usr/bin/env python3
"""Build separate table, figure-legend, and cover-letter submission files."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/jijiabiao/Desktop/耳聋课题/2026_supplement_analysis/manuscript")
SOURCE = ROOT / "manuscript.md"
OUT = ROOT / "submission_package"

BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
HEADER_FILL = "1F4E79"
ALT_FILL = "EAF2F8"
BORDER = "D9D9D9"


def set_run_font(run, size: float, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(12)

    for style_name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    add_page_number(section.footer.paragraphs[0])


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)
    set_run_font(run, 9)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def parse_markdown() -> tuple[list[str], list[tuple[str, list[list[str]]]]]:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    legends: list[str] = []
    tables: list[tuple[str, list[list[str]]]] = []

    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith("**Figure "):
            match = re.match(r"\*\*(.*?)\*\*\s*(.*)", line)
            if match:
                legends.append(f"{match.group(1)} {match.group(2)}".strip())
        if line.startswith("**Table ") and line.endswith("**"):
            caption = line.strip("*")
            i += 1
            while i < len(lines) and not lines[i].strip():
                i += 1
            table_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for table_line in table_lines:
                cells = [cell.strip() for cell in table_line.strip("|").split("|")]
                if all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
                    continue
                rows.append(cells)
            tables.append((caption, rows))
            continue
        i += 1
    return legends, tables


def add_table(doc: Document, rows: list[list[str]], column_widths: list[float] | None = None) -> None:
    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for column_index in range(column_count):
            value = row[column_index] if column_index < len(row) else ""
            cell = cells[column_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(value)
            set_run_font(run, 9.5, row_index == 0)
            if row_index == 0:
                set_cell_shading(cell, HEADER_FILL)
                run.font.color.rgb = WHITE
            elif row_index % 2 == 0:
                set_cell_shading(cell, ALT_FILL)
    if column_widths:
        for row in table.rows:
            for cell, width in zip(row.cells, column_widths):
                cell.width = Inches(width)
    else:
        width = 6.5 / column_count
        for row in table.rows:
            for cell in row.cells:
                cell.width = Inches(width)
    if table.rows:
        repeat_table_header(table.rows[0])
    doc.add_paragraph()


def build_figure_legends(legends: list[str]) -> None:
    doc = Document()
    configure_document(doc)
    title = doc.add_paragraph(style="Title")
    set_run_font(title.add_run("Figure Legends"), 20, True)
    subtitle = doc.add_paragraph()
    set_run_font(
        subtitle.add_run("Peripheral-Central Convergence of Macrophage and Microglial Programs in Presbycusis Severity"),
        11,
        True,
    )
    for legend in legends:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        set_run_font(paragraph.add_run(legend), 11)
    output = OUT / "Figure_Legends.docx"
    doc.save(output)
    print(f"Wrote {output}")


def build_tables(tables: list[tuple[str, list[list[str]]]]) -> None:
    doc = Document()
    configure_document(doc)
    title = doc.add_paragraph(style="Title")
    set_run_font(title.add_run("Main Tables 1-6"), 20, True)
    subtitle = doc.add_paragraph()
    set_run_font(
        subtitle.add_run("Peripheral-Central Convergence of Macrophage and Microglial Programs in Presbycusis Severity"),
        11,
        True,
    )
    for index, (caption, rows) in enumerate(tables):
        if index:
            doc.add_page_break()
        heading = doc.add_heading(caption, level=1)
        for run in heading.runs:
            set_run_font(run, 12.5, True)
        add_table(doc, rows)
    output = OUT / "Tables" / "Main_Tables_1-6.docx"
    doc.save(output)
    print(f"Wrote {output}")


def build_cover_letter() -> None:
    doc = Document()
    configure_document(doc)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.08
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    doc.styles["Title"].font.size = Pt(18)
    title = doc.add_paragraph(style="Title")
    set_run_font(title.add_run("Cover Letter"), 20, True)

    doc.add_paragraph(date.today().strftime("%B %-d, %Y"))
    doc.add_paragraph("Editorial Office\nHearing Research")
    doc.add_paragraph("Dear Editors,")

    paragraphs = [
        (
            "We are pleased to submit our manuscript, \"Peripheral-Central Convergence of Macrophage and "
            "Microglial Programs in Presbycusis Severity,\" for consideration as an Original Research Article. "
            "The study presents an integrative transcriptomic analysis of the aging mouse cochlea and inferior "
            "colliculus using severity-resolved modeling rather than a binary disease comparison."
        ),
        (
            "Presbycusis is clinically heterogeneous and involves both peripheral and central auditory dysfunction. "
            "Most prior transcriptomic studies have focused on the cochlea and used single-tissue hub-gene or "
            "binary group comparisons. Our study addresses this gap by integrating GSE49543 and GSE49522, performing "
            "limma/eBayes severity modeling and WGCNA, comparing peripheral and central auditory tissues, and localizing "
            "candidate programs with single-nucleus RNA sequencing from GSE274279."
        ),
        (
            "The main findings are that middle age alone did not produce robust differential expression, whereas "
            "progression toward severe presbycusis was associated with a macrophage and microglial program. A "
            "66-gene WGCNA module was strongly correlated with severity and preserved in the inferior colliculus. "
            "Among shared peripheral-central genes, 48 were significant in both tissues and were enriched for immune "
            "system processes, macrophage biology, and microglial biology. Single-nucleus data localized the leading "
            "candidate genes primarily to macrophages/microglia. External datasets showed partial directional "
            "reproducibility, and nested cross-validation demonstrated that internal model performance did not transfer "
            "to an external cohort. We therefore present the machine-learning analysis only as a secondary, "
            "hypothesis-generating analysis."
        ),
        (
            "An exhaustive CellChat analysis of all 2,019 CellChatDB.mouse interactions showed a monotonic decline in "
            "cochlear communication with age, a preserved core of 46 ligand-receptor pairs, PTN signaling as the "
            "strongest pathway at every age, macrophage/microglial PTPRC-MRC1 communication at all ages, and APP-CD74 "
            "signaling that became significant at 12 and 24 months. Transcription-factor enrichment (ChEA), miRNA "
            "target enrichment (TargetScan), drug-signature enrichment (DSigDB), and GO and KEGG enrichment converged "
            "on macrophage-associated regulators and on phagocytic and antigen-presentation programs. The miRNA "
            "analysis did not reach FDR significance and is reported as a negative result."
        ),
        (
            "We believe the manuscript fits the scope of Hearing Research because it directly addresses presbycusis "
            "severity, cochlear aging, central auditory involvement, and immune cell biology in the auditory system. "
            "The study is a secondary analysis of publicly available mouse data. No new human participants, animal "
            "experiments, or identifiable patient data were used. The work has not been published previously and is "
            "not under consideration elsewhere. All authors have approved the manuscript and agree with its submission. "
            "The authors declare no competing interests."
        ),
        (
            "The manuscript, separate figure files, main tables, Figure Legends, and Supplementary Tables S1-S37 are "
            "provided with this submission. We would be grateful for your consideration."
        ),
    ]
    for text in paragraphs:
        paragraph = doc.add_paragraph(text)
        paragraph.paragraph_format.space_after = Pt(8)

    doc.add_paragraph("Sincerely,")
    doc.add_paragraph(
        "Jianming Yang, on behalf of all authors\n"
        "The Second Affiliated Hospital of Anhui Medical University\n"
        "Hefei, Anhui, China\n"
        "Email and telephone: to be confirmed before submission"
    )
    output = OUT / "Cover_Letter_Hearing_Research.docx"
    doc.save(output)
    print(f"Wrote {output}")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "Tables").mkdir(parents=True, exist_ok=True)
    legends, tables = parse_markdown()
    build_figure_legends(legends)
    build_tables(tables)
    build_cover_letter()


if __name__ == "__main__":
    main()
