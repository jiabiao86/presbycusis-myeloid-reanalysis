#!/usr/bin/env python3
"""Build a submission-ready DOCX draft from the frozen manuscript Markdown."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
HEADER_FILL = "1F4E79"
ALT_FILL = "EAF2F8"
BORDER = "D9D9D9"


def set_run_font(run, name: str, size: float, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold


def add_inline_text(paragraph, text: str, size: float, bold: bool = False) -> None:
    """Render simple Markdown emphasis without changing scientific notation."""
    pattern = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, "Arial", size, bold)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, "Arial", size, True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Arial", size, bold)
            run.italic = True
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, "Arial", size, bold)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, "Arial", 9)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(14)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for style_name, size in (("Heading 1", 14.5), ("Heading 2", 12.5), ("Heading 3", 11.5)):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 9)
        style.paragraph_format.space_after = Pt(5)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    add_page_number(footer_paragraph)

    props = doc.core_properties
    props.title = "Peripheral-Central Convergence of Macrophage and Microglial Programs in Presbycusis Severity"
    props.subject = "Integrative transcriptomic analysis of aging mouse cochlea and inferior colliculus"
    props.author = "Manuscript draft"
    props.keywords = "presbycusis; macrophage; microglia; cochlea; inferior colliculus; WGCNA"


def parse_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)
    return rows


def set_table_column_widths(table, rows: list[list[str]]) -> None:
    width = max(len(row) for row in rows)
    header = rows[0]
    if width == 2:
        widths = [1.8, 5.0]
    elif width == 3:
        widths = [1.5, 2.4, 2.9]
    elif width == 5 and header and header[0] == "Comparison":
        widths = [1.55, 1.25, 1.25, 1.25, 1.25]
    elif width == 5 and header and header[0] == "Gene":
        widths = [0.65, 0.85, 0.85, 1.15, 3.3]
    elif width == 5 and header and header[0] == "Dataset":
        widths = [0.8, 1.2, 0.95, 1.6, 2.25]
    elif width == 7 and "Example genes" in header:
        widths = [0.55, 0.55, 0.8, 0.85, 0.6, 0.85, 2.4]
    else:
        widths = [6.8 / width] * width

    for row in table.rows:
        for cell, width_in in zip(row.cells, widths):
            cell.width = Inches(width_in)


def add_table(doc: Document, rows: list[list[str]], caption_line: str | None) -> None:
    if caption_line:
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)
        add_inline_text(p, caption_line, 9.5, True)

    width = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col_index in range(width):
            value = row[col_index] if col_index < len(row) else ""
            cell = cells[col_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline_text(paragraph, value, 9.0, row_index == 0)
            if row_index == 0:
                set_cell_shading(cell, HEADER_FILL)
                for run in paragraph.runs:
                    run.font.color.rgb = WHITE
            elif row_index % 2 == 0:
                set_cell_shading(cell, ALT_FILL)
        if row_index == 0:
            repeat_table_header(table.rows[0])

    set_table_column_widths(table, rows)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(5)


def build(markdown_path: Path, output_path: Path) -> None:
    text = markdown_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    configure_document(doc)

    base_dir = markdown_path.parent
    i = 0
    pending_caption: str | None = None
    title_seen = False
    keyword_seen = False
    figure_caption_count = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue

        image_match = re.fullmatch(r"!\[[^\]]*\]\(([^)]+)\)", line)
        if image_match:
            image_path = base_dir / image_match.group(1)
            image_alt = re.match(r"!\[([^\]]*)\]", line).group(1)
            doc.add_picture(str(image_path), width=Inches(6.55))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for doc_pr in doc.paragraphs[-1]._p.xpath(".//wp:docPr"):
                doc_pr.set("descr", image_alt)
                doc_pr.set("title", image_alt)
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(5)
            i += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, parse_table(table_lines), pending_caption)
            pending_caption = None
            continue

        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            add_inline_text(p, line[2:].strip(), 20, True)
            title_seen = True
        elif line.startswith("## "):
            heading_text = line[3:].strip()
            if heading_text in {"Figures", "Tables"}:
                doc.add_page_break()
            p = doc.add_paragraph(style="Heading 1")
            add_inline_text(p, heading_text, 14.5, True)
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline_text(p, line[4:].strip(), 12.5, True)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            add_inline_text(p, line[2:].strip(), 11)
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.space_after = Pt(4)
            add_inline_text(p, line, 10.2)
        elif line.startswith("**Figure "):
            figure_caption_count += 1
            if figure_caption_count > 1:
                doc.add_page_break()
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            add_inline_text(p, line, 10.2, True)
        elif line.startswith("**Table ") and line.endswith("**"):
            if line.startswith("**Table 6."):
                doc.add_page_break()
            pending_caption = line
        else:
            p = doc.add_paragraph()
            add_inline_text(p, line, 11)
            if line.startswith("**Keywords:**"):
                keyword_seen = True
            if title_seen and keyword_seen and line.startswith("**Keywords:**"):
                p.add_run().add_break(WD_BREAK.PAGE)
        i += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    build(args.input, args.output)


if __name__ == "__main__":
    main()
